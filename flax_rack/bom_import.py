"""BOM .xlsx -> geometry seed + contradiction report (spec 5.3).

Best effort for the SIMS F4 BOM sheet family ("F4 rack elevation" +
"Cable Guide v5"); not a general parser. The returned dict is a SEED: rack tag,
mgmt-switch uplink ports and iperf server are left for the operator to fix by
hand. openpyxl is imported lazily (bang host code is otherwise stdlib only).
"""
import re
import warnings

PARTS = {"compute-tp": "801-0023", "bc-head-tp": "801-0025", "bc-jbod": "802-0010"}

_TITLES = [
    (re.compile(r"^Compute Tioga Pass ([LCR])(\d\d)$"), "compute-tp"),
    (re.compile(r"^Bryce Canyon head Tioga Pass ([LCR])(\d\d)$"), "bc-head-tp"),
    (re.compile(r"^Bryce Canyon JBOD (\d\d)\b"), "bc-jbod"),
    (re.compile(r"^Celestica switch\s*-\s*([A-Z])$"), "switch-mgmt"),
    (re.compile(r"^Wedge 100\s*-\s*([A-Z])$"), "switch-tor"),
]
_SW_ROLE = {"switch-mgmt": "mgmt", "switch-tor": "tor"}
_PCIE = re.compile(r"^PCIe P(\d)\s*:\s*SW ([A-Z]) P(\d+)$")
_SWPORT = re.compile(r"^SW ([A-Z]) P(\d+)$")
_NIC = re.compile(r"^NIC P(\d)$")
_LANE = re.compile(r"^[A-D]([1-4])$")
_STOP = re.compile(r"^(Empty|Powershelf)$")
_COLRANGE = {"L": (2, 9), "C": (10, 17), "R": (18, 25)}  # sheet columns B-I, J-Q, R-Y


def _txt(v):
    return re.sub(r"\s+", " ", str(v).replace("\xa0", " ")).strip() if v is not None else ""


def _sw_model(title):
    """The model as the BOM names it ("Celestica switch - C" -> "celestica",
    "Wedge 100 - A" -> "wedge100"). A swapped-in replacement is a hand edit."""
    return re.sub(r"[^a-z0-9]", "", re.sub(r"\bswitch\b", "", title.rsplit("-", 1)[0].lower()))


def _cable_guide(ws):
    """{(col L/C/R, ou, nic 0/1): (wedge, port, lane)} from the breakout tables."""
    out = {}
    for lab in (50, 58, 66):  # AX, BF, BN: label column of Left / Center / Right
        wedge = port = None
        for r in range(1, ws.max_row + 1):
            t = _txt(ws.cell(r, lab).value)
            m = re.match(r"^Wedge ([AB])$", t)
            if m:
                p = re.match(r"^port (\d+)$", _txt(ws.cell(r, lab + 4).value))
                wedge, port = m.group(1), int(p.group(1)) if p else None
            elif t == "Breakout" and wedge:
                lane = _LANE.match(_txt(ws.cell(r, lab + 2).value))
                col = _txt(ws.cell(r, lab + 3).value)
                ou = _txt(ws.cell(r, lab + 4).value)
                nic = re.match(r"^NIC P(\d)$", _txt(ws.cell(r, lab + 5).value))
                if lane and col in _COLRANGE and ou.isdigit() and nic:
                    # a breakout row names the lane/OU; port comes from the header above
                    out[(col, int(ou), int(nic.group(1)))] = (wedge, port, int(lane.group(1)))
    return out


def import_bom(xlsx_path, *, order, customer, site, vid, prefix):
    import openpyxl
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")  # pivot-cache UserWarning
        wb = openpyxl.load_workbook(str(xlsx_path), data_only=True)
    ws = wb["F4 rack elevation"]
    guide = _cable_guide(wb["Cable Guide v5"]) if "Cable Guide v5" in wb.sheetnames else {}

    def ou_label(r):
        v = _txt(ws.cell(r, 1).value)
        return int(v) if v.isdigit() else None

    titles = []  # (row, col, kind, groups)
    for r in range(1, ws.max_row + 1):
        for c in range(2, 26):
            t = _txt(ws.cell(r, c).value)
            for rx, kind in _TITLES:
                m = rx.match(t)
                if m:
                    titles.append((r, c, kind, m.groups(), t))

    switches, devices, contradictions = [], [], []
    for r, c, kind, g, title in titles:
        top = ou_label(r)
        if top is None:
            continue
        if kind.startswith("switch"):
            switches.append({"id": g[0], "role": _SW_ROLE[kind], "ou": top,
                             "model": _sw_model(title),
                             "nos": "sonic", "mgmt_ip": None,
                             "uplink_ports": []})
            continue
        # height: OU label rows from the title down until an Empty/Powershelf/next title
        h = 1
        rr = r + 2
        while ou_label(rr) is not None:
            cells = [_txt(ws.cell(rr, cc).value) for cc in range(2, 26)]
            if any(_STOP.match(t) or any(rx.match(t) for rx, _ in _TITLES) for t in cells):
                break
            h += 1
            rr += 2
        bottom = top - h + 1
        if kind == "bc-jbod":
            pos, col = f"JBOD{g[0]}", None
            halves = {"iom1": (2, 13), "iom2": (14, 25)}
        else:
            col = g[0]
            pos = f"{col}{g[1]}"
            lo, hi = _COLRANGE[col]
            halves = {None: (lo, hi)}
        links, lane = {}, None
        for rr in range(r + 1, r + 2 * h):
            for cc in range(2, 26):
                t = _txt(ws.cell(rr, cc).value)
                if not t:
                    continue
                for key, (lo, hi) in halves.items():
                    if not lo <= cc <= hi:
                        continue
                    m = _LANE.match(t)
                    if m and lane is None:
                        lane = int(m.group(1))
                    m = _PCIE.match(t)
                    if m:
                        links[f"pcie{m.group(1)}"] = (m.group(2), int(m.group(3)))
                    m = _NIC.match(t)
                    s = _SWPORT.match(_txt(ws.cell(rr + 1, cc).value)) if m else None
                    if s:
                        name = key or f"mezz{m.group(1)}"
                        links[name] = (s.group(1), int(s.group(2)))
        out = {}
        for k, (sw, port) in links.items():
            out[k] = f"{sw}:{port}/{lane}" if k.startswith("pcie") and lane else f"{sw}:{port}"
        dev = {"pos": pos, "ou": bottom, "h": h, "col": col, "kind": kind,
               "part": PARTS[kind], "links": out}
        devices.append(dev)

        if col and kind in ("compute-tp", "bc-head-tp"):
            for n in (0, 1):
                k = f"pcie{n}"
                if k not in links:
                    continue
                # exact (col, OU); else the guide's topmost entry with the same lane
                # (the guide predates the OU37 computes and starts at its D4 slot)
                gk = guide.get((col, bottom, n)) or guide.get((col, top, n))
                if gk is None and lane is not None:
                    cands = [(o, v) for (cc_, o, nn), v in guide.items()
                             if cc_ == col and nn == n and v[2] == lane]
                    gk = max(cands)[1] if cands else None
                if gk is None:
                    continue
                gsw, gport, glane = gk
                if (gsw, gport) != links[k]:
                    contradictions.append(
                        f"{pos} {k}: elevation {links[k][0]}:{links[k][1]}/{lane} "
                        f"vs cable table {gsw}:{gport}/{glane}")

    switches.sort(key=lambda s: -s["ou"])
    devices.sort(key=lambda d: -d["ou"])
    geometry = {
        "order": order, "rack": "", "customer": customer, "site": site,
        "network": {"vid": vid, "cidr": f"172.{vid}.0.0/16", "prefix": prefix},
        "iperf_server": None, "switches": switches, "devices": devices,
    }
    return geometry, contradictions
